"""Gộp noi_dung_3.3.md, noi_dung_3.4_linh_kien.md, noi_dung_4.4.md và noi_dung_phu_luc_A.md thành một file .docx nháp.

Chỉ lấy phần trước dòng '---' của mỗi file (phần sau là ghi chú nội bộ của đội).
Hỗ trợ tập con Markdown đang dùng: tiêu đề '#', đoạn văn, **đậm**, *nghiêng*, `mã`, ảnh + chú thích,
bảng '|', danh sách '- ' và '1. '.
Chạy: python3 _scripts/tao_docx.py
"""
from __future__ import annotations

import io
import re
import sys
from pathlib import Path

from PIL import Image

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCES = [
    ROOT / "3.3_kien_truc" / "noi_dung_3.3.md",
    ROOT / "3.4_linh_kien" / "noi_dung_3.4_linh_kien.md",
    ROOT / "4.4_hinh_minh_hoa" / "noi_dung_4.4.md",
]
# Phụ lục: bản vẽ mạch khổ ngang, đặt trong section A4 nằm ngang để chữ trên bản vẽ đủ lớn.
APPENDIX = ROOT / "3.3_kien_truc" / "noi_dung_phu_luc_A.md"
IMG_CM, IMG_CM_NGANG = 16, 24     # 24 cm: ảnh tỉ lệ A3 + chú thích vừa trang A4 ngang
OUT = ROOT / "EdgeFlow_ban_nhap_muc_3.3_3.4_4.4.docx"
FONT = "Times New Roman"
MONO = "Consolas"
BOM_WIDTHS_CM = [1.1, 3.6, 5.1, 2.0, 4.2]
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def set_font(run, name: str = FONT, size: float | None = None) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)
    if size:
        run.font.size = Pt(size)


def setup_styles(doc: Document) -> None:
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(13)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.25
    for name, size in (("Heading 1", 16), ("Heading 2", 14)):
        h = doc.styles[name]
        h.font.name = FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        h.element.rPr.rFonts.set(qn("w:ascii"), FONT)
        h.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)


def add_inline(par, text: str, size: float | None = None, italic: bool = False) -> None:
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2])
            r.bold = True
            set_font(r, size=size)
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1])
            set_font(r, MONO, (size or 13) - 1.5)
        elif tok.startswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            r.italic = True
            set_font(r, size=size)
        else:
            r = par.add_run(tok)
            set_font(r, size=size)
        if italic:
            r.italic = True


def shade(cell, hex_fill: str) -> None:
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    ncol = len(rows[0])
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = BOM_WIDTHS_CM if ncol == 5 else [16.0 / ncol] * ncol
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        group = i > 0 and row[0].startswith("**") and not any(c.strip() for c in row[2:])
        for j, txt in enumerate(row):
            c = cells[j]
            c.width = Cm(widths[j])
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_inline(p, txt.strip(), size=10.5)
            if i == 0:
                for r in p.runs:
                    r.bold = True
                shade(c, "E7E7E7")
            elif group:
                shade(c, "F3F3F3")
        if group:
            merged = cells[1].merge(cells[ncol - 1])
            merged.paragraphs[0].runs[0].bold = True
    doc.add_paragraph()


def embed_bytes(path: Path) -> io.BytesIO:
    """Ảnh nhúng vào Word: render → JPEG 1900 px (300 dpi ở 16 cm), sơ đồ → PNG ≤ 2400 px."""
    im = Image.open(path)
    is_render = path.name.startswith(("H4.", "H3.5"))
    max_w = 1900 if is_render else 2400
    if im.width > max_w:
        im = im.resize((max_w, round(im.height * max_w / im.width)), Image.LANCZOS)
    buf = io.BytesIO()
    if is_render:
        im.convert("RGB").save(buf, "JPEG", quality=88, optimize=True)
    else:
        im.convert("RGB").save(buf, "PNG", optimize=True)
    buf.seek(0)
    return buf


def add_image(doc: Document, path: Path, caption: str | None, width_cm: float = IMG_CM) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(embed_bytes(path), width=Cm(width_cm))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(cp, caption, size=11, italic=True)


def add_landscape_section(doc: Document) -> None:
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.5)


def convert(doc: Document, md_path: Path, img_cm: float = IMG_CM) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.strip() == "---":
            break
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=2)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=3)
        elif line.startswith("!["):
            m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", line)
            cap = None
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            if j < len(lines) and re.fullmatch(r"\*[^*].*\*", lines[j].strip()):
                cap = lines[j].strip()[1:-1]
                i = j
            add_image(doc, md_path.parent / m.group(1), cap, img_cm)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not re.fullmatch(r"[\s:\-|]+", lines[i]):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        elif line.startswith("- "):
            add_inline(doc.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"\d+\. ", line):
            # số viết sẵn: tránh việc Word đánh số nối tiếp giữa các danh sách
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, line)
        else:
            p = doc.add_paragraph()
            if line.startswith("**Bảng"):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
            add_inline(p, line)
        i += 1


def main() -> int:
    doc = Document()
    setup_styles(doc)
    t = doc.add_paragraph()
    add_inline(t, "**EdgeFlow AI (đội 31µF): bản nháp mục 3.3, 3.4 (linh kiện) và 4.4**", size=15)
    note = doc.add_paragraph()
    add_inline(note, "Nháp ngày 11/10/2026 để dán vào mẫu báo cáo bán kết. Hình 3.1–3.4 lấy từ file draw.io; "
                     "Hình 3.5 và 4.1–4.6 là render 3D minh họa; Hình A.1–A.4 (Phụ lục A, khổ ngang) "
                     "lấy từ file draw.io bản vẽ mạch.", size=11, italic=True)
    for src in SOURCES:
        convert(doc, src)
    add_landscape_section(doc)
    convert(doc, APPENDIX, IMG_CM_NGANG)
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")  # template mặc định của python-docx thiếu thuộc tính bắt buộc này
    doc.save(OUT)
    print("->", OUT.relative_to(ROOT))
    return 0


if __name__ == "__main__":
    sys.exit(main())
